"""
file_import.py — parse an uploaded CSV/XLSX of locations + niches into a
ParsedImport. No persistence, no side effects. .xls (old binary format) is
not supported.
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

_LOCATION_HEADERS = ("city", "town", "location", "city/town")
_STATE_HEADERS = ("state", "province", "region", "st")
# "type" was removed: it is the weakest niche signal and matched a City/Town
# location-classification column, turning a places list into niches=["City","Town"].
_NICHE_HEADERS = ("niche", "industry", "category", "keyword", "service")

# A "niche" column filled entirely with these is a location classification, not
# a list of business categories — reject the upload rather than build a queue
# that searches for the literal word "City" in every town.
_PLACE_TYPE_WORDS = frozenset({
    "city", "town", "village", "hamlet", "county", "state", "province",
    "municipality", "borough", "township", "district", "region", "area",
    "place", "locality", "suburb", "neighborhood", "neighbourhood",
})


@dataclass
class ParsedImport:
    locations: List[Dict[str, Optional[str]]] = field(default_factory=list)
    niches: List[str] = field(default_factory=list)
    combinations: int = 0
    layout: str = "combined"          # "combined" | "separate"
    warnings: List[str] = field(default_factory=list)


def _norm(h: Any) -> str:
    return str(h or "").strip().lower().replace("_", " ").replace("-", " ")


def _match(header: Any, candidates: Tuple[str, ...]) -> bool:
    h = _norm(header)
    if not h:
        return False
    tokens = set(h.replace("/", " ").split())
    # Exact or whole-token match only — a bare `c in h` substring test wrongly
    # matched "st" inside "First Name", "town" inside "Downtown", etc.
    return any(c == h or c in tokens for c in candidates)


def _reject_or_warn_place_type_niches(niches: List[str], warnings: List[str]) -> None:
    """Guard against a location list being imported as the niche list. If every
    niche value is a place type (City / Town / County …) the column is a
    location classification — reject it. If only some are, warn and continue."""
    if not niches:
        return
    place_like = sorted({n for n in niches if n.strip().lower() in _PLACE_TYPE_WORDS})
    if len(place_like) == len({n.strip().lower() for n in niches}):
        raise ValueError(
            "The niche/category column contains place types "
            f"({', '.join(place_like)}), not business categories. Use a column "
            "of niches such as 'Dental Clinic', 'Law Firm', 'Restaurant'."
        )
    if place_like:
        warnings.append(
            f"{len(place_like)} niche value(s) look like place types "
            f"({', '.join(place_like)}) and may not be valid business categories."
        )


def _rows_from_csv(data: bytes) -> List[List[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    return [row for row in csv.reader(io.StringIO(text))]


def _sheets_from_xlsx(data: bytes) -> Dict[str, List[List[Any]]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: Dict[str, List[List[Any]]] = {}
    for ws in wb.worksheets:
        out[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return out


def _extract_table(
    rows: List[List[Any]],
) -> Tuple[Optional[int], Optional[int], Optional[int], List[List[str]]]:
    """Returns (city_idx, state_idx, niche_idx, body_rows). Indices are None if
    the column is absent. body_rows are the non-empty rows as trimmed strings."""
    if not rows:
        return None, None, None, []
    header = rows[0]
    city_idx = state_idx = niche_idx = None
    for i, h in enumerate(header):
        if city_idx is None and _match(h, _LOCATION_HEADERS):
            city_idx = i
        elif state_idx is None and _match(h, _STATE_HEADERS):
            state_idx = i
        elif niche_idx is None and _match(h, _NICHE_HEADERS):
            niche_idx = i
    body = []
    for r in rows[1:]:
        cells = [("" if c is None else str(c)).strip() for c in r]
        if any(cells):
            body.append(cells)
    return city_idx, state_idx, niche_idx, body


def _dedupe(seq: List[str]) -> Tuple[List[str], int]:
    seen, out, dups = set(), [], 0
    for s in seq:
        k = s.lower()
        if k in seen:
            dups += 1
        else:
            seen.add(k)
            out.append(s)
    return out, dups


def _dedupe_locations(
    locs: List[Dict[str, Optional[str]]],
) -> Tuple[List[Dict[str, Optional[str]]], int]:
    seen, out, dups = set(), [], 0
    for loc in locs:
        k = (loc["city"].lower(), (loc["state"] or "").lower())
        if k in seen:
            dups += 1
        else:
            seen.add(k)
            out.append(loc)
    return out, dups


def _load(name: str, blob: bytes) -> Dict[str, List[List[Any]]]:
    low = (name or "").lower()
    if low.endswith(".csv"):
        return {"__csv__": _rows_from_csv(blob)}
    if low.endswith(".xlsx"):
        return _sheets_from_xlsx(blob)
    raise ValueError(f"Unsupported file type '{name}'. Supported: .csv, .xlsx")


def _locs_from(body: List[List[str]], ci: int, si: Optional[int]) -> List[Dict[str, Optional[str]]]:
    out = []
    for r in body:
        if ci < len(r) and r[ci]:
            state = r[si] if (si is not None and si < len(r)) else None
            out.append({"city": r[ci], "state": state or None})
    return out


def parse_upload(
    filename: str, data: bytes,
    second_filename: Optional[str] = None, second_data: Optional[bytes] = None,
) -> ParsedImport:
    sheets: Dict[str, List[List[Any]]] = _load(filename, data)
    if second_data is not None:
        for k, v in _load(second_filename or "second.csv", second_data).items():
            sheets[f"{k}#2"] = v

    warnings: List[str] = []

    # ── Combined: a single table with both a location and a niche column ──
    for _name, rows in sheets.items():
        ci, si, ni, body = _extract_table(rows)
        if ci is not None and ni is not None:
            locs = _locs_from(body, ci, si)
            niches = [r[ni] for r in body if ni < len(r) and r[ni]]
            locs, dl = _dedupe_locations(locs)
            niches, dn = _dedupe(niches)
            if dl or dn:
                warnings.append(f"Collapsed {dl + dn} duplicate row(s).")
            _reject_or_warn_place_type_niches(niches, warnings)
            if not locs or not niches:
                raise ValueError("File parsed but produced no usable locations or niches.")
            return ParsedImport(locs, niches, len(locs) * len(niches), "combined", warnings)

    # ── Separate: one table has locations only, another has niches only ──
    loc_rows: List[Dict[str, Optional[str]]] = []
    niche_vals: List[str] = []
    for _name, rows in sheets.items():
        ci, si, ni, body = _extract_table(rows)
        if ci is not None and ni is None:
            loc_rows += _locs_from(body, ci, si)
        elif ni is not None:
            niche_vals += [r[ni] for r in body if ni < len(r) and r[ni]]
        elif ci is None and ni is None and rows:
            # A bare single-column sheet with no recognised header — treat it
            # as a hand-made niche list (header cell included as data unless it
            # is itself a niche-header word).
            width = max((len(r) for r in rows), default=0)
            if width == 1:
                header_cell = str(rows[0][0]).strip() if rows[0] else ""
                vals = ([header_cell] if header_cell and not _match(header_cell, _NICHE_HEADERS) else [])
                vals += [str(r[0]).strip() for r in rows[1:] if r and r[0] is not None]
                niche_vals += [v for v in vals if v]

    loc_rows, dl = _dedupe_locations(loc_rows)
    niche_vals, dn = _dedupe(niche_vals)
    if dl or dn:
        warnings.append(f"Collapsed {dl + dn} duplicate row(s).")
    _reject_or_warn_place_type_niches(niche_vals, warnings)
    if not niche_vals:
        raise ValueError("No niche column found. Expected a column named niche / industry / category.")
    if not loc_rows:
        raise ValueError("No location column found. Expected a column named city / town / location.")
    return ParsedImport(loc_rows, niche_vals, len(loc_rows) * len(niche_vals), "separate", warnings)
