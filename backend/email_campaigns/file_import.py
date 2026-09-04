"""
file_import.py — parse an uploaded lead file (CSV / XLSX) into normalized lead
rows for an Email Campaign. Pure function, no persistence, no side effects.

Adopts the n8n V1.1 "Normalize Lead" header-mapping so arbitrary column names
(First Name / firstname / fname → first_name, etc.) resolve to canonical
fields. Preserves the original row verbatim as `raw`. Preserves a supplied
`body` column BYTE-FOR-BYTE (only leading/trailing whitespace is trimmed, to
match n8n V1.1 and the storage layer).

.xls (old binary format) is NOT parsed here — the same limitation as
backend/automation/file_import.py (no xlrd dependency). .xls files are handled
by the n8n preparation step (Checkpoint 3C) instead; import_leads reports that
clearly.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional

from ..validators import is_valid_email

# canonical field  ->  accepted header spellings (already lower-cased + space-normalized)
_HEADER_MAP: Dict[str, List[str]] = {
    "first_name":      ["first name", "firstname", "first", "fname", "given name"],
    "last_name":       ["last name", "lastname", "last", "lname", "surname", "family name"],
    "email":           ["email", "email address", "e mail", "mail", "emailaddress"],
    "company":         ["company", "company name", "organization", "organisation", "business",
                        "business name", "account"],
    "job_title":       ["job title", "jobtitle", "title", "position", "role"],
    "website":         ["website", "site", "url", "web site", "domain"],
    "phone":           ["phone", "phone number", "tel", "telephone", "mobile", "cell"],
    "industry":        ["industry", "sector", "niche", "vertical"],
    "location":        ["location", "city", "town"],
    "personalization": ["personalization", "personalisation", "note", "notes", "custom", "hook"],
    "body":            ["body", "email body", "email_body", "message", "message body",
                        "email content", "content", "custom body"],
}
_CANON_FIELDS = list(_HEADER_MAP.keys())

_MAX_ROWS = 50_000


@dataclass
class ImportedLead:
    lead_key:      str
    email:         Optional[str]
    first_name:    str = ""
    last_name:     str = ""
    company:       str = ""
    body_source:   str = "ai"          # 'ai' | 'provided'
    provided_body: Optional[str] = None
    status:        str = "VALIDATED"    # VALIDATED / MISSING_EMAIL / INVALID_EMAIL / DUPLICATE
    status_detail: str = ""
    raw:           Dict[str, Any] = field(default_factory=dict)


@dataclass
class ImportResult:
    leads:        List[ImportedLead] = field(default_factory=list)
    total_rows:   int = 0
    valid:        int = 0
    missing_email: int = 0
    invalid_email: int = 0
    duplicates:   int = 0
    warnings:     List[str] = field(default_factory=list)


def _norm_header(h: Any) -> str:
    return re.sub(r"[\s_\-]+", " ", str(h or "").strip().lower()).strip()


def _canon(header: Any) -> Optional[str]:
    nh = _norm_header(header)
    for canon, spellings in _HEADER_MAP.items():
        if nh == canon or nh in spellings:
            return canon
    return None


def _slug(*parts: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", "-".join(p for p in parts if p).lower()).strip("-") or "unknown"


def _rows_from_csv(data: bytes) -> List[List[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    return [row for row in csv.reader(io.StringIO(text))]


def _rows_from_xlsx(data: bytes) -> List[List[Any]]:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]  # first sheet only
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def is_xls_binary(filename: str) -> bool:
    return (filename or "").lower().endswith(".xls")


def parse_lead_file(campaign_id: int, filename: str, data: bytes) -> ImportResult:
    low = (filename or "").lower()
    if low.endswith(".csv"):
        rows = _rows_from_csv(data)
    elif low.endswith(".xlsx"):
        rows = _rows_from_xlsx(data)
    elif low.endswith(".xls"):
        raise ValueError(
            ".xls (old binary Excel) is not supported by the native importer — "
            "re-save as .xlsx or .csv, or run the campaign through n8n preparation."
        )
    else:
        raise ValueError(f"Unsupported file type: {filename!r}. Supported: .csv, .xlsx")

    if not rows:
        raise ValueError("File is empty.")
    header = rows[0]
    if len(rows) - 1 > _MAX_ROWS:
        raise ValueError(f"Too many rows ({len(rows) - 1}); limit is {_MAX_ROWS}.")

    # header index -> canonical field (unmapped headers kept in raw only)
    col_map: Dict[int, str] = {}
    for i, h in enumerate(header):
        c = _canon(h)
        if c and c not in col_map.values():
            col_map[i] = c

    result = ImportResult()
    seen_keys: set[str] = set()
    seen_emails: set[str] = set()

    for rownum, r in enumerate(rows[1:], start=1):
        cells = ["" if c is None else str(c) for c in r]
        raw = {str(header[i] if i < len(header) else f"col{i}"): cells[i] for i in range(len(cells))}
        if not any(v.strip() for v in cells):
            continue
        result.total_rows += 1

        vals: Dict[str, str] = {}
        for i, canon in col_map.items():
            vals[canon] = cells[i].strip() if i < len(cells) else ""

        email_raw = vals.get("email", "").strip()
        email_lc = email_raw.lower()
        body_raw = (vals.get("body") or "")
        provided = body_raw.strip()  # trim only — matches n8n V1.1 + storage layer

        lead = ImportedLead(
            lead_key="",  # set below
            email=email_lc or None,
            first_name=vals.get("first_name", ""),
            last_name=vals.get("last_name", ""),
            company=vals.get("company", ""),
            body_source="provided" if provided else "ai",
            provided_body=provided or None,
            raw=raw,
        )

        # lead_key: <campaign>::<email>  or  <campaign>::noemail::<slug>::<row>
        if email_lc:
            lead.lead_key = f"{campaign_id}::{email_lc}"
        else:
            lead.lead_key = f"{campaign_id}::noemail::{_slug(lead.first_name, lead.last_name, lead.company)}::{rownum}"

        # ---- validation + dedupe ----
        if not email_raw:
            lead.status, lead.status_detail = "MISSING_EMAIL", "email is missing"
            result.missing_email += 1
        elif not is_valid_email(email_lc):
            lead.status, lead.status_detail = "INVALID_EMAIL", "email format is invalid"
            result.invalid_email += 1
        elif lead.lead_key in seen_keys or (email_lc and email_lc in seen_emails):
            lead.status, lead.status_detail = "DUPLICATE", "duplicate within this file"
            result.duplicates += 1
        else:
            lead.status = "VALIDATED"
            result.valid += 1
            seen_keys.add(lead.lead_key)
            if email_lc:
                seen_emails.add(email_lc)

        result.leads.append(lead)

    if result.total_rows == 0:
        raise ValueError("File parsed but contained no data rows.")
    if not col_map or "email" not in col_map.values():
        result.warnings.append("No 'email' column recognised — every lead will be MISSING_EMAIL.")

    return result
